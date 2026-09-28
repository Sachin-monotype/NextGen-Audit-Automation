"""Notification test catalog + per-cycle actuals (local JSON, Excel export).

Catalog is seeded from ``python/audit_validator/data/notification_test_catalog.json``
(imported from Final Test Notifications.xlsx). Cycle results live under
``reports/notification-cycles/{env}/{cycle_id}.json`` so each test pass / env
keeps its own Actual / Status / Comments without rewriting the source sheet.
"""

from __future__ import annotations

import json
import re
import threading
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

_lock = threading.Lock()
_ENVS = frozenset({"pp", "qa", "uat", "beta"})


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _catalog_path(project_root: Path) -> Path:
    return project_root / "python" / "audit_validator" / "data" / "notification_test_catalog.json"


def _cycles_root(project_root: Path) -> Path:
    return project_root / "reports" / "notification-cycles"


def _slug(text: str) -> str:
    raw = re.sub(r"[^a-zA-Z0-9._-]+", "-", (text or "").strip().lower()).strip("-")
    return raw or "cycle"


def load_catalog(project_root: Path) -> dict[str, Any]:
    path = _catalog_path(project_root)
    if not path.is_file():
        return {
            "version": 1,
            "title": "Notification Test Catalog",
            "channels": ["in_app", "email", "push"],
            "statuses": ["", "Pass", "Fail", "Blocked", "N/A"],
            "categories": [],
            "items": [],
            "error": f"Catalog missing: {path}",
        }
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        return data if isinstance(data, dict) else {"items": [], "categories": []}
    except Exception as exc:
        return {"items": [], "categories": [], "error": str(exc)}


def list_cycles(project_root: Path, *, env: str | None = None) -> list[dict[str, Any]]:
    root = _cycles_root(project_root)
    if not root.is_dir():
        return []
    out: list[dict[str, Any]] = []
    envs = [env] if env and env in _ENVS else sorted(_ENVS)
    for e in envs:
        folder = root / e
        if not folder.is_dir():
            continue
        for path in sorted(folder.glob("*.json"), key=lambda p: p.stat().st_mtime, reverse=True):
            try:
                data = json.loads(path.read_text(encoding="utf-8"))
            except Exception:
                continue
            if not isinstance(data, dict):
                continue
            results = data.get("results") if isinstance(data.get("results"), dict) else {}
            filled = sum(
                1
                for v in results.values()
                if isinstance(v, dict)
                and (
                    str(v.get("status") or "").strip()
                    or str(v.get("actual") or "").strip()
                    or str(v.get("comments") or "").strip()
                )
            )
            out.append(
                {
                    "id": data.get("id") or path.stem,
                    "name": data.get("name") or path.stem,
                    "env": data.get("env") or e,
                    "created_at": data.get("created_at") or "",
                    "updated_at": data.get("updated_at") or "",
                    "item_count": len(results),
                    "filled_count": filled,
                    "path": str(path.relative_to(project_root)),
                }
            )
    return out


def _cycle_path(project_root: Path, env: str, cycle_id: str) -> Path:
    e = (env or "uat").strip().lower()
    if e not in _ENVS:
        e = "uat"
    return _cycles_root(project_root) / e / f"{_slug(cycle_id)}.json"


def load_cycle(project_root: Path, env: str, cycle_id: str) -> dict[str, Any] | None:
    path = _cycle_path(project_root, env, cycle_id)
    if not path.is_file():
        return None
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        return data if isinstance(data, dict) else None
    except Exception:
        return None


def create_cycle(
    project_root: Path,
    *,
    env: str,
    name: str,
    seed_from_catalog: bool = True,
) -> dict[str, Any]:
    e = (env or "uat").strip().lower()
    if e not in _ENVS:
        e = "uat"
    label = (name or "").strip() or f"{e}-cycle"
    cycle_id = f"{_slug(label)}-{datetime.now(timezone.utc).strftime('%Y%m%d-%H%M%S')}"
    results: dict[str, Any] = {}
    if seed_from_catalog:
        catalog = load_catalog(project_root)
        for item in catalog.get("items") or []:
            if not isinstance(item, dict):
                continue
            iid = str(item.get("id") or "").strip()
            if not iid:
                continue
            results[iid] = {
                "status": str(item.get("seed_status") or "").strip(),
                "actual": str(item.get("seed_actual") or "").strip(),
                "comments": str(item.get("seed_comments") or "").strip(),
                "how_to_notes": "",
                "tested_at": "",
                "tester": "",
            }
    doc = {
        "id": cycle_id,
        "name": label,
        "env": e,
        "created_at": _now(),
        "updated_at": _now(),
        "results": results,
    }
    path = _cycle_path(project_root, e, cycle_id)
    path.parent.mkdir(parents=True, exist_ok=True)
    with _lock:
        path.write_text(json.dumps(doc, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return doc


def upsert_result(
    project_root: Path,
    *,
    env: str,
    cycle_id: str,
    item_id: str,
    patch: dict[str, Any],
) -> dict[str, Any]:
    path = _cycle_path(project_root, env, cycle_id)
    with _lock:
        if path.is_file():
            try:
                doc = json.loads(path.read_text(encoding="utf-8"))
            except Exception:
                doc = {}
        else:
            doc = {
                "id": cycle_id,
                "name": cycle_id,
                "env": (env or "uat").strip().lower() or "uat",
                "created_at": _now(),
                "results": {},
            }
        if not isinstance(doc, dict):
            doc = {"id": cycle_id, "results": {}}
        results = doc.get("results") if isinstance(doc.get("results"), dict) else {}
        current = results.get(item_id) if isinstance(results.get(item_id), dict) else {}
        allowed = ("status", "actual", "comments", "how_to_notes", "tester", "tested_at")
        merged = dict(current)
        for key in allowed:
            if key in patch:
                merged[key] = patch[key]
        if any(k in patch for k in ("status", "actual", "comments")) and not merged.get("tested_at"):
            merged["tested_at"] = _now()
        results[str(item_id)] = merged
        doc["results"] = results
        doc["updated_at"] = _now()
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(doc, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        return doc


def merged_view(project_root: Path, *, env: str, cycle_id: str | None) -> dict[str, Any]:
    catalog = load_catalog(project_root)
    cycle = load_cycle(project_root, env, cycle_id) if cycle_id else None
    results = (cycle or {}).get("results") if isinstance((cycle or {}).get("results"), dict) else {}
    items_out: list[dict[str, Any]] = []
    for item in catalog.get("items") or []:
        if not isinstance(item, dict):
            continue
        iid = str(item.get("id") or "")
        res = results.get(iid) if isinstance(results.get(iid), dict) else {}
        channels = item.get("channels") if isinstance(item.get("channels"), dict) else {}
        items_out.append(
            {
                "id": iid,
                "category_id": item.get("category_id") or "",
                "category": item.get("category") or "",
                "permission": item.get("permission") or "",
                "recipients": item.get("recipients") or "",
                "event": item.get("event") or "",
                "trigger": item.get("trigger") or "",
                "how_to": item.get("how_to") or item.get("trigger") or "",
                "channels": {
                    "in_app": bool(channels.get("in_app")),
                    "email": bool(channels.get("email")),
                    "push": bool(channels.get("push")),
                },
                "expected": item.get("expected") or "",
                "status": res.get("status") or "",
                "actual": res.get("actual") or "",
                "comments": res.get("comments") or "",
                "how_to_notes": res.get("how_to_notes") or "",
                "tester": res.get("tester") or "",
                "tested_at": res.get("tested_at") or "",
            }
        )
    return {
        "catalog": {
            "title": catalog.get("title") or "Notification Test Catalog",
            "categories": catalog.get("categories") or [],
            "statuses": catalog.get("statuses")
            or ["", "Pass", "Fail", "Blocked", "N/A"],
            "channels": catalog.get("channels") or ["in_app", "email", "push"],
            "source": catalog.get("source") or "",
            "error": catalog.get("error"),
        },
        "cycle": {
            "id": (cycle or {}).get("id"),
            "name": (cycle or {}).get("name"),
            "env": (cycle or {}).get("env") or env,
            "created_at": (cycle or {}).get("created_at"),
            "updated_at": (cycle or {}).get("updated_at"),
        }
        if cycle
        else None,
        "items": items_out,
        "cycles": list_cycles(project_root, env=env),
    }


def export_excel_bytes(project_root: Path, *, env: str, cycle_id: str) -> bytes:
    view = merged_view(project_root, env=env, cycle_id=cycle_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "Notifications"
    headers = [
        "#",
        "Category",
        "Event",
        "Trigger / How to",
        "In-App",
        "Email",
        "Push",
        "Status",
        "Expected Notifications",
        "Actual Notifications",
        "Comments",
        "Permission",
        "Recipients",
        "Tester",
        "Tested at",
    ]
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    for col, h in enumerate(headers, 1):
        cell = ws.cell(1, col, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    def mark(v: bool) -> str:
        return "✓" if v else "—"

    for row_i, item in enumerate(view.get("items") or [], 2):
        ch = item.get("channels") or {}
        values = [
            item.get("id") or "",
            item.get("category") or "",
            item.get("event") or "",
            item.get("trigger") or "",
            mark(bool(ch.get("in_app"))),
            mark(bool(ch.get("email"))),
            mark(bool(ch.get("push"))),
            item.get("status") or "",
            item.get("expected") or "",
            item.get("actual") or "",
            item.get("comments") or "",
            item.get("permission") or "",
            item.get("recipients") or "",
            item.get("tester") or "",
            item.get("tested_at") or "",
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row_i, col, val)
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    from openpyxl.utils import get_column_letter

    widths = [8, 28, 36, 40, 8, 8, 8, 10, 48, 48, 28, 28, 28, 14, 22]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    cycle = view.get("cycle") or {}
    meta = wb.create_sheet("Cycle")
    meta["A1"] = "Environment"
    meta["B1"] = cycle.get("env") or env
    meta["A2"] = "Cycle"
    meta["B2"] = cycle.get("name") or cycle_id
    meta["A3"] = "Exported at"
    meta["B3"] = _now()

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
